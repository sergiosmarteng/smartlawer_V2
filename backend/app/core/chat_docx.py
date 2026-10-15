"""Geração em memória de arquivo .docx para respostas do chat jurídico.

Converte marcações Markdown emitidas pelo modelo (títulos `#`/`##`/`###`,
listas com marcador ou numeradas, citações `>` e ênfases `**negrito**` /
`*itálico*`) em parágrafos e estilos nativos do Word (`python-docx`), sem
deixar sintaxe Markdown crua no documento final e sem gravar arquivos
temporários em disco.
"""

from __future__ import annotations

import re
from datetime import datetime, timezone
from io import BytesIO
from typing import Iterable

from docx import Document as DocxDocument
from docx.shared import Pt, RGBColor

from app.schemas.chat import Citation

_HEADING_RE = re.compile(r"^\s*(#{1,6})\s+(.+?)\s*#*$")
_BULLET_RE = re.compile(r"^\s*[-*•]\s+(.+)$")
_NUMBERED_RE = re.compile(r"^\s*(\d+)[.)]\s+(.+)$")
_QUOTE_RE = re.compile(r"^\s*>\s*(.+)$")
_HR_RE = re.compile(r"^\s*(?:---+|\*\*\*+|___+)\s*$")

# Captura tokens inline: **negrito**, __negrito__, `código`, *itálico*
_INLINE_TOKEN_RE = re.compile(
    r"(\*\*[^*\n]+\*\*|__[^_\n]+__|`[^`\n]+`|\*[^*\n]+\*)"
)


def _append_formatted_runs(paragraph, text: str) -> None:
    """Adiciona runs ao parágrafo convertendo **negrito**, *itálico* e `código`."""
    cleaned = text.replace("\r", "")
    pos = 0
    for match in _INLINE_TOKEN_RE.finditer(cleaned):
        start, end = match.span()
        if start > pos:
            paragraph.add_run(cleaned[pos:start].replace("**", ""))
        token = match.group(0)
        if (token.startswith("**") and token.endswith("**")) or (
            token.startswith("__") and token.endswith("__")
        ):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
        elif token.startswith("*") and token.endswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        else:
            paragraph.add_run(token)
        pos = end
    if pos < len(cleaned):
        paragraph.add_run(cleaned[pos:].replace("**", ""))


def _strip_heading_inline_markers(text: str) -> str:
    """Remove delimitadores ** ou * de títulos mantendo apenas o texto limpo."""
    text = re.sub(r"\*\*([^*\n]+)\*\*", r"\1", text)
    text = re.sub(r"__([^_\n]+)__", r"\1", text)
    text = re.sub(r"\*([^*\n]+)\*", r"\1", text)
    text = re.sub(r"`([^`\n]+)`", r"\1", text)
    return text.replace("**", "").strip()


def _render_answer_blocks(doc: DocxDocument, answer: str) -> None:
    lines = (answer or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")
    paragraph_buffer: list[str] = []

    def flush_paragraph() -> None:
        if not paragraph_buffer:
            return
        joined = " ".join(part.strip() for part in paragraph_buffer if part.strip())
        paragraph_buffer.clear()
        if not joined:
            return
        p = doc.add_paragraph()
        _append_formatted_runs(p, joined)

    for raw_line in lines:
        line = raw_line.strip()
        if not line:
            flush_paragraph()
            continue

        if _HR_RE.match(line):
            flush_paragraph()
            continue

        heading_match = _HEADING_RE.match(raw_line)
        if heading_match:
            flush_paragraph()
            hashes, title_text = heading_match.groups()
            level = min(max(len(hashes), 1), 3) + 1  # H1->2, H2->3, H3+->4
            doc.add_heading(_strip_heading_inline_markers(title_text), level=level)
            continue

        bullet_match = _BULLET_RE.match(raw_line)
        if bullet_match:
            flush_paragraph()
            p = doc.add_paragraph(style="List Bullet")
            _append_formatted_runs(p, bullet_match.group(1).strip())
            continue

        numbered_match = _NUMBERED_RE.match(raw_line)
        if numbered_match:
            flush_paragraph()
            p = doc.add_paragraph(style="List Number")
            _append_formatted_runs(p, numbered_match.group(2).strip())
            continue

        quote_match = _QUOTE_RE.match(raw_line)
        if quote_match:
            flush_paragraph()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            _append_formatted_runs(p, quote_match.group(1).strip())
            for run in p.runs:
                run.italic = True
            continue

        paragraph_buffer.append(line)

    flush_paragraph()


def build_chat_answer_docx(
    *,
    question: str,
    answer: str,
    document_name: str | None = None,
    citations: Iterable[Citation] = (),
) -> bytes:
    """Constrói um arquivo .docx em memória para uma resposta do chat jurídico."""
    doc = DocxDocument()

    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(11)

    title = doc.add_heading("SmartLawer — Parecer do Chat Jurídico", level=1)
    for run in title.runs:
        run.font.color.rgb = RGBColor(0x1E, 0x2E, 0x23)

    now_str = datetime.now(timezone.utc).strftime("%d/%m/%Y %H:%M UTC")
    scope_label = document_name or "Todos os documentos do histórico"
    meta_p = doc.add_paragraph()
    meta_run = meta_p.add_run(f"Caso / escopo: {scope_label}  •  Gerado em: {now_str}")
    meta_run.italic = True
    meta_run.font.size = Pt(9.5)
    meta_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    clean_question = (question or "").strip()
    if clean_question:
        doc.add_heading("Pergunta", level=2)
        q_p = doc.add_paragraph()
        _append_formatted_runs(q_p, clean_question)

    doc.add_heading("Resposta fundamentada", level=2)
    _render_answer_blocks(doc, answer)

    citations_list = list(citations or [])
    if citations_list:
        doc.add_heading("Fontes citadas", level=2)
        for item in citations_list:
            page_suffix = (
                f" · p. {item.page_start}"
                if isinstance(item.page_start, int)
                else ""
            )
            cite_p = doc.add_paragraph(style="List Bullet")
            header_run = cite_p.add_run(
                f"{item.ref} {item.document_name}{page_suffix}"
            )
            header_run.bold = True
            excerpt = (item.excerpt or "").strip()
            if excerpt:
                cite_p.add_run(f" — {excerpt}")

    doc.add_heading("Aviso de revisão profissional (HITL)", level=2)
    disclaimer_p = doc.add_paragraph()
    disclaimer_run = disclaimer_p.add_run(
        "Rascunho elaborado por inteligência artificial (ai_draft=true, "
        "requires_human_review=true). Todo o conteúdo, citações e conclusões "
        "devem ser revisados e validados pelo advogado responsável antes de "
        "qualquer uso ou protocolo."
    )
    disclaimer_run.italic = True
    disclaimer_run.font.size = Pt(9.5)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
