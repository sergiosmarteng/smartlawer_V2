"""Testes para exportação .docx de respostas individuais do chat jurídico."""

from io import BytesIO

from docx import Document as DocxDocument

from app.models.audit_event import AuditEvent
from app.models.document import Document


def test_chat_export_docx_requires_auth(client):
    resp = client.post(
        "/api/v1/chat/export-docx",
        json={"question": "Qual a tese?", "answer": "Resposta fundamentada."},
    )
    assert resp.status_code == 401


def test_chat_export_docx_generates_structured_docx_without_raw_markdown(
    client, db_session, make_user, auth_headers_for
):
    user = make_user()
    document = Document(
        user_id=user.id,
        filename="Contestacao_Silva.pdf",
        file_path="/tmp/contestacao.pdf",
        content_type="application/pdf",
        status="completed",
    )
    db_session.add(document)
    db_session.commit()

    markdown_answer = (
        "### 1. Síntese da tese\n\n"
        "A cobrança é **indevida** conforme o contrato [1] e a análise [A].\n\n"
        "- Primeiro ponto com **destaque** relevante\n"
        "- Segundo ponto com *observação* complementar\n\n"
        "1. Requerer juntada do contrato\n"
        "2. Impugnar planilha de cálculo\n"
    )
    payload = {
        "question": "Qual a tese de defesa mais forte deste caso?",
        "answer": markdown_answer,
        "document_id": str(document.id),
        "citations": [
            {
                "ref": "[1]",
                "chunk_id": "c-1",
                "document_id": str(document.id),
                "document_name": "Contestacao_Silva.pdf",
                "page_start": 4,
                "excerpt": "Cláusula 5ª prevê isenção expressa.",
            }
        ],
    }

    resp = client.post(
        "/api/v1/chat/export-docx",
        headers=auth_headers_for(user),
        json=payload,
    )
    assert resp.status_code == 200
    assert (
        resp.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert "attachment;" in resp.headers.get("content-disposition", "")
    assert ".docx" in resp.headers.get("content-disposition", "")
    assert resp.content[:2] == b"PK"

    parsed = DocxDocument(BytesIO(resp.content))
    full_text = "\n".join(p.text for p in parsed.paragraphs if p.text.strip())

    # Conteúdo essencial presente
    assert "Qual a tese de defesa mais forte deste caso?" in full_text
    assert "Contestacao_Silva.pdf" in full_text
    assert "1. Síntese da tese" in full_text
    assert "Primeiro ponto com destaque relevante" in full_text
    assert "Requerer juntada do contrato" in full_text
    assert "Cláusula 5ª prevê isenção expressa." in full_text
    assert "revisão" in full_text.lower() or "advogado" in full_text.lower()

    # Sem sintaxe Markdown crua no texto final do Word
    assert "###" not in full_text
    assert "**indevida**" not in full_text
    assert "**destaque**" not in full_text

    # Verifica que negrito foi aplicado como run real do Word
    bold_runs = [
        run.text
        for p in parsed.paragraphs
        for run in p.runs
        if run.bold and run.text.strip()
    ]
    assert any("indevida" in r or "destaque" in r for r in bold_runs)

    # Auditoria registrada sem vazar pergunta ou resposta (LGPD)
    event = (
        db_session.query(AuditEvent)
        .filter(
            AuditEvent.user_id == user.id,
            AuditEvent.event_type == AuditEvent.DOCX_GENERATED,
        )
        .order_by(AuditEvent.created_at.desc())
        .first()
    )
    assert event is not None
    assert event.entity_id == str(document.id)
    meta_dump = str(event.meta or {})
    assert "indevida" not in meta_dump
    assert "Qual a tese" not in meta_dump
    assert (event.meta or {}).get("source") == "chat"


def test_chat_export_docx_rejects_other_users_document(
    client, db_session, make_user, auth_headers_for
):
    owner = make_user(email="owner-chat-docx@example.com", username="owner_cd")
    intruder = make_user(email="intruder-chat-docx@example.com", username="intruder_cd")
    document = Document(
        user_id=owner.id,
        filename="Sigiloso.pdf",
        file_path="/tmp/sigiloso.pdf",
        content_type="application/pdf",
        status="completed",
    )
    db_session.add(document)
    db_session.commit()

    resp = client.post(
        "/api/v1/chat/export-docx",
        headers=auth_headers_for(intruder),
        json={
            "question": "Pergunta",
            "answer": "Resposta",
            "document_id": str(document.id),
        },
    )
    assert resp.status_code == 404
